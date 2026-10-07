from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

# Initialize presentation
prs = Presentation()
prs.slide_width = Inches(13.333)  # 16:9 Widescreen
prs.slide_height = Inches(7.5)

# Slide Data for Section 1 (5 Chunks)
slides_data = [
    {
        "slide_num": "SLIDE 1 (0:00 - 0:03)",
        "voiceover": '"Most CS grads fail their tech interviews—not because code fails, but because it looks textbook."',
        "visual": "Split screen: Left side Crimson Red (#2A0808) with Red X icon. Right side Emerald Green (#032010) with Green Checkmark icon.",
        "assets": '• transparent red X icon png\n• transparent green checkmark png',
        "effects": 'Glitch Transition on opening. Typewriter Text on top: "JUNIOR VS SENIOR"',
        "sfx": "• Harsh Cyberpunk Glitch / Digital Static at 0:00\n• Background Track: Cyberpunk Tech Synth (10% Vol)"
    },
    {
        "slide_num": "SLIDE 2 (0:03 - 0:06)",
        "voiceover": '"If a senior hiring manager sees you handling nested objects with endless IF statements..."',
        "visual": "Zoom into Left (Red) Side. Display screenshot of a messy VS Code screen with 5 nested if statements.",
        "assets": '• VS Code screenshot with nested if statements\n• Red highlight box overlay',
        "effects": "Crop to Fill / Zoom. Apply Flash Red Color Filter for 1 second.",
        "sfx": "• Low Error Buzz / Vinyl Scratch sound at 0:04"
    },
    {
        "slide_num": "SLIDE 3 (0:06 - 0:09)",
        "voiceover": '"...it instantly signals you’ve never built software in production."',
        "visual": 'Overlay a giant Red Stamp across the screen saying: "REJECTED".',
        "assets": '• transparent rubber stamp rejected red png',
        "effects": 'Slow Zoom In. Text Overlay: "NO PRODUCTION EXPERIENCE"',
        "sfx": "• Heavy Metal Slam / Thud sound when stamp hits screen"
    },
    {
        "slide_num": "SLIDE 4 (0:09 - 0:12)",
        "voiceover": '"In the next 5 minutes, I\'m going to show you the single line of modern JavaScript..."',
        "visual": "Screen wipes to Right (Green) Side. Display a clean 1-line code snippet glowing in green.",
        "assets": '• neon green glow background overlay free',
        "effects": 'Wipe Right Transition. Text Overlay: "1-LINE PRODUCTION HACK"',
        "sfx": "• Fast Whoosh Sound Effect on wipe"
    },
    {
        "slide_num": "SLIDE 5 (0:12 - 0:15)",
        "voiceover": '"...that senior engineers use to eliminate nested checks, stop crashes, and get hired."',
        "visual": "Pop-up Graphic: Green Offer Letter or Tech Corporate Logo badges (FAANG style).",
        "assets": '• job offer letter icon green png',
        "effects": 'Bounce Animation to offer letter graphic. Text Overlay: "HIRED ✅"',
        "sfx": "• Clean Digital Chime / Level Up Sound"
    }
]

blank_layout = prs.slide_layouts[6]

for data in slides_data:
    slide = prs.slides.add_slide(blank_layout)
    
    # Background Dark Gradient Feel
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor(18, 18, 24)
    
    # Slide Title Box
    title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(11.7), Inches(0.8))
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    p.text = data["slide_num"]
    p.font.bold = True
    p.font.size = Pt(28)
    p.font.color.rgb = RGBColor(255, 230, 0)  # Neon Yellow
    
    # Left Content Box (Voiceover + Visual)
    left_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.5), Inches(5.6), Inches(5.2))
    tf_l = left_box.text_frame
    tf_l.word_wrap = True
    
    p1 = tf_l.paragraphs[0]
    p1.text = "🎙️ Voiceover (Audio):"
    p1.font.bold = True
    p1.font.size = Pt(18)
    p1.font.color.rgb = RGBColor(0, 230, 150)
    
    p2 = tf_l.add_paragraph()
    p2.text = data["voiceover"] + "\n"
    p2.font.size = Pt(15)
    p2.font.color.rgb = RGBColor(255, 255, 255)
    
    p3 = tf_l.add_paragraph()
    p3.text = "🖼️ Visual Elements & Setup:"
    p3.font.bold = True
    p3.font.size = Pt(18)
    p3.font.color.rgb = RGBColor(0, 180, 255)
    
    p4 = tf_l.add_paragraph()
    p4.text = data["visual"]
    p4.font.size = Pt(15)
    p4.font.color.rgb = RGBColor(220, 220, 220)

    # Right Content Box (Assets, Effects, Audio)
    right_box = slide.shapes.add_textbox(Inches(6.8), Inches(1.5), Inches(5.7), Inches(5.2))
    tf_r = right_box.text_frame
    tf_r.word_wrap = True
    
    r1 = tf_r.paragraphs[0]
    r1.text = "🔍 Google Asset Searches:"
    r1.font.bold = True
    r1.font.size = Pt(18)
    r1.font.color.rgb = RGBColor(255, 180, 0)
    
    r2 = tf_r.add_paragraph()
    r2.text = data["assets"] + "\n"
    r2.font.size = Pt(14)
    r2.font.color.rgb = RGBColor(220, 220, 220)
    
    r3 = tf_r.add_paragraph()
    r3.text = "✨ Clipchamp Effects & Text Overlays:"
    r3.font.bold = True
    r3.font.size = Pt(18)
    r3.font.color.rgb = RGBColor(200, 100, 255)
    
    r4 = tf_r.add_paragraph()
    r4.text = data["effects"] + "\n"
    r4.font.size = Pt(14)
    r4.font.color.rgb = RGBColor(220, 220, 220)

    r5 = tf_r.add_paragraph()
    r5.text = "🎵 SFX & Audio Cues:"
    r5.font.bold = True
    r5.font.size = Pt(18)
    r5.font.color.rgb = RGBColor(255, 80, 80)
    
    r6 = tf_r.add_paragraph()
    r6.text = data["sfx"]
    r6.font.size = Pt(14)
    r6.font.color.rgb = RGBColor(220, 220, 220)

# Save the presentation
prs.save("Section_1_Editing_Storyboard.pptx")
print("Successfully generated editable presentation: Section_1_Editing_Storyboard.pptx")